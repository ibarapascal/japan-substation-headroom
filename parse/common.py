"""Shared low-level parsing helpers: text decoding, zip member names, cell values, timestamps,
voltages, and row readers for CSV / XLSX / XLS sources.

Conventions
- Missing-value tokens published by the operators ("-", "ー", "－", "―", "‐", "欠測", "＊", "*",
  empty cell) are all read as missing (None).
- Thousands separators are removed ("1,800" -> 1800). Full-width digits and minus signs are
  normalised by NFKC before conversion.
- A leading apostrophe (used by one operator to force text cells) and surrounding quotes are
  stripped from labels and values.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import math
import re
import unicodedata
import zipfile
from typing import Iterator, List, Optional, Sequence, Tuple

MISSING_TOKENS = {"", "-", "ー", "－", "―", "‐", "‑", "–", "—", "−", "─", "欠測", "＊", "*",
                  "#N/A", "N/A", "NA", "nan", "NaN", "None"}


def nfkc(s) -> str:
    return unicodedata.normalize("NFKC", "" if s is None else str(s))


def clean_label(s) -> str:
    """Normalise a header label or text cell: NFKC, strip quotes/apostrophes/whitespace."""
    t = nfkc(s).strip()
    t = t.strip("'\"` \t\r\n　")
    return t


def decode_text(raw: bytes) -> Tuple[str, str]:
    """Decode bytes as UTF-8 (with or without BOM), falling back to CP932."""
    for enc in ("utf-8-sig", "cp932"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("cp932", errors="replace"), "cp932-replace"


def zip_member_name(info: zipfile.ZipInfo) -> str:
    """Zip entries without the UTF-8 flag are stored as CP932 bytes; Python reads them as CP437."""
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp932")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def iter_zip(path_or_bytes) -> Iterator[Tuple[str, bytes]]:
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes
    with zipfile.ZipFile(src) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            yield zip_member_name(info), z.read(info)


def parse_number(v) -> Optional[float]:
    """Cell value -> float, or None for missing/non-numeric."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if math.isnan(f) else f
    s = clean_label(v)
    if s in MISSING_TOKENS:
        return None
    s = s.replace(",", "").replace("−", "-").replace("–", "-")
    try:
        return float(s)
    except ValueError:
        return None


_KV_NUM = re.compile(r"\d+(?:\.\d+)?")


def norm_kv(v) -> str:
    """Voltage cell -> canonical string: '66', '6.6', '22-13.8', '110/66'. '' if missing."""
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and math.isnan(v):
            return ""
        return _fmt_num(float(v))
    s = clean_label(v)
    if s in MISSING_TOKENS:
        return ""
    s = re.sub(r"(?i)kv", "", s).strip()
    nums = _KV_NUM.findall(s)
    if not nums:
        return ""
    if len(nums) == 1:
        return _fmt_num(float(nums[0]))
    sep = "-" if "-" in s else "/"
    return sep.join(_fmt_num(float(n)) for n in nums)


def _fmt_num(f: float) -> str:
    return str(int(f)) if f == int(f) else ("%g" % f)


def kv_join_key(kv: str) -> str:
    """Voltage string -> join key; 6/6.6/6.9 kV are one class, 3/3.3 kV another."""
    if not kv:
        return ""
    try:
        f = float(kv)
    except ValueError:
        return kv
    if 6.0 <= f < 7.0:
        return "6.6"
    if 3.0 <= f < 3.5:
        return "3.3"
    return kv


def kv_value(kv: str) -> Optional[float]:
    """Largest numeric value in a voltage string (used only for classification)."""
    nums = _KV_NUM.findall(kv or "")
    return max(float(n) for n in nums) if nums else None


# --- timestamps -------------------------------------------------------------------------------

_TS_PATTERNS = [
    re.compile(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?$"),
    # two-digit year with hour only, e.g. "23/04/01 01"
    re.compile(r"^(\d{2})/(\d{1,2})/(\d{1,2})\s+(\d{1,2})()$"),
    re.compile(r"^(\d{4})年(\d{1,2})月(\d{1,2})日\s*(\d{1,2})時(?:(\d{1,2})分(?:\d{1,2}秒)?)?$"),
]
_DATE_PATTERNS = [
    re.compile(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?:\s+0{1,2}:00(?::00)?)?$"),
    re.compile(r"^(\d{4})年(\d{1,2})月(\d{1,2})日$"),
]
_HOUR_PATTERN = re.compile(r"^(\d{1,2})(?::(\d{2})(?::\d{2})?|時)?$")


def parse_ts(v) -> Optional[dt.datetime]:
    """Cell -> naive local datetime. Hour 24 rolls over to 00 of the next day."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.replace(tzinfo=None, microsecond=0)
    s = clean_label(v)
    if not s:
        return None
    for pat in _TS_PATTERNS:
        m = pat.match(s)
        if m:
            y, mo, d, h = (int(m.group(i)) for i in range(1, 5))
            if y < 100:
                y += 2000
            mi = int(m.group(5) or 0)
            try:
                base = dt.datetime(y, mo, d)
            except ValueError:
                return None
            if h > 24 or mi > 59:
                return None
            return base + dt.timedelta(hours=h, minutes=mi)
    return None


def parse_date(v) -> Optional[dt.date]:
    if isinstance(v, dt.datetime):
        return v.date() if (v.hour, v.minute) == (0, 0) else None
    if isinstance(v, dt.date):
        return v
    s = clean_label(v)
    for pat in _DATE_PATTERNS:
        m = pat.match(s)
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


def parse_hour(v) -> Optional[int]:
    """'1:00', '1', '01時', time object, or a float fraction of a day -> hour 0..24."""
    if v is None:
        return None
    if isinstance(v, dt.time):
        return v.hour
    if isinstance(v, dt.timedelta):
        return int(round(v.total_seconds() / 3600))
    if isinstance(v, float) and 0 <= v < 1.0001 and v != int(v):
        return int(round(v * 24))
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v) if 0 <= v <= 24 and v == int(v) else None
    m = _HOUR_PATTERN.match(clean_label(v))
    if m:
        h = int(m.group(1))
        return h if 0 <= h <= 24 else None
    return None


def fiscal_year(t: dt.datetime) -> int:
    return t.year if t.month >= 4 else t.year - 1


def fy_start(fy: int) -> dt.datetime:
    return dt.datetime(fy, 4, 1)


def fy_hours(fy: int) -> int:
    return int((fy_start(fy + 1) - fy_start(fy)).total_seconds() // 3600)


# --- row readers ------------------------------------------------------------------------------

Row = List[object]


def csv_rows(raw: bytes) -> Tuple[Iterator[Row], str]:
    """Rows of the decoded text via csv.reader, lazily (never splitlines: cells may contain
    embedded newlines and commas)."""
    text, enc = decode_text(raw)
    return csv.reader(io.StringIO(text)), enc


def xlsx_sheets(raw: bytes) -> Iterator[Tuple[str, Iterator[Sequence[object]]]]:
    import openpyxl  # optional dependency
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title, ws.iter_rows(values_only=True)
    finally:
        wb.close()


def xls_sheets(raw: bytes) -> Iterator[Tuple[str, Iterator[Sequence[object]]]]:
    import xlrd  # optional dependency
    wb = xlrd.open_workbook(file_contents=raw, on_demand=True)
    for sh in wb.sheets():
        def rows(sh=sh):
            for i in range(sh.nrows):
                out = []
                for j in range(sh.ncols):
                    c = sh.cell(i, j)
                    if c.ctype == xlrd.XL_CELL_DATE:
                        try:
                            out.append(xlrd.xldate_as_datetime(c.value, wb.datemode))
                        except Exception:
                            out.append(c.value)
                    elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        out.append(None)
                    else:
                        out.append(c.value)
                yield out
        yield sh.name, rows()
